
import os
import sqlite3
from datetime import datetime, date
from pathlib import Path

from aiogram import Bot, Dispatcher, F
from aiogram.filters import Command, CommandStart
from aiogram.types import Message, CallbackQuery
from aiogram.utils.keyboard import InlineKeyboardBuilder
from openpyxl import load_workbook

TOKEN = os.getenv("BOT_TOKEN")
DB_PATH = os.getenv("DB_PATH", "tracker.db")
PLAN_PATH = os.getenv("PLAN_PATH", "plan.xlsx")

if not TOKEN:
    raise RuntimeError("BOT_TOKEN is not set")

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS tasks (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        task_date TEXT NOT NULL,
        day TEXT,
        load TEXT,
        task TEXT NOT NULL,
        skill TEXT,
        minutes INTEGER DEFAULT 0,
        notes TEXT
    );
    CREATE TABLE IF NOT EXISTS users (
        telegram_id INTEGER PRIMARY KEY,
        first_name TEXT
    );
    CREATE TABLE IF NOT EXISTS completions (
        telegram_id INTEGER NOT NULL,
        task_id INTEGER NOT NULL,
        completed_at TEXT NOT NULL,
        PRIMARY KEY (telegram_id, task_id)
    );
    """)
    conn.commit()
    conn.close()

def import_plan():
    conn = db()
    count = conn.execute("SELECT COUNT(*) FROM tasks").fetchone()[0]
    if count:
        conn.close()
        return

    wb = load_workbook(PLAN_PATH, data_only=True)
    ws = wb["Daily Tracker"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    idx = {h:i for i,h in enumerate(headers)}

    for r in rows[1:]:
        if not r[idx["Date"]] or not r[idx["Task"]]:
            continue
        d = r[idx["Date"]]
        if hasattr(d, "strftime"):
            d = d.strftime("%Y-%m-%d")
        else:
            d = str(d)[:10]
        minutes = r[idx["Minutes"]] or 0
        try:
            minutes = int(float(minutes))
        except:
            minutes = 0
        conn.execute(
            "INSERT INTO tasks(task_date,day,load,task,skill,minutes,notes) VALUES(?,?,?,?,?,?,?)",
            (d, r[idx["Day"]], r[idx["Load"]], r[idx["Task"]], r[idx["Skill"]], minutes, r[idx["Notes"]])
        )
    conn.commit()
    conn.close()

def save_user(message: Message):
    conn=db()
    conn.execute(
        "INSERT INTO users(telegram_id,first_name) VALUES(?,?) "
        "ON CONFLICT(telegram_id) DO UPDATE SET first_name=excluded.first_name",
        (message.from_user.id, message.from_user.first_name or "")
    )
    conn.commit(); conn.close()

def get_tasks(d: str, uid: int):
    conn=db()
    rows=conn.execute("""
      SELECT t.*, CASE WHEN c.task_id IS NULL THEN 0 ELSE 1 END AS done
      FROM tasks t LEFT JOIN completions c
      ON t.id=c.task_id AND c.telegram_id=?
      WHERE t.task_date=? ORDER BY t.id
    """,(uid,d)).fetchall()
    conn.close()
    return rows

def task_keyboard(task_id, done):
    b=InlineKeyboardBuilder()
    b.button(
        text="↩️ Mark as not done" if done else "✅ Done",
        callback_data=f"toggle:{task_id}"
    )
    b.adjust(1)
    return b.as_markup()

def fmt_task(i, row):
    status="✅" if row["done"] else "⬜"
    return f"{status} <b>{row['skill']}</b> — {row['minutes']} min\n{row['task']}"

async def send_today(message: Message, target_date=None):
    uid=message.from_user.id
    d=target_date or date.today().isoformat()
    rows=get_tasks(d,uid)
    if not rows:
        await message.answer(f"📅 <b>{d}</b>\n\nBu sana uchun task topilmadi.")
        return

    done=sum(r["done"] for r in rows)
    total_minutes=sum(r["minutes"] for r in rows)
    done_minutes=sum(r["minutes"] for r in rows if r["done"])
    load=rows[0]["load"] or ""

    await message.answer(
        f"📅 <b>{rows[0]['day']} — {d}</b>\n"
        f"⚡ Load: <b>{load}</b>\n"
        f"🎯 Progress: <b>{done}/{len(rows)}</b> ({round(done/len(rows)*100)}%)\n"
        f"⏱ Time: <b>{done_minutes}/{total_minutes} min</b>"
    )
    for i,row in enumerate(rows,1):
        await message.answer(fmt_task(i,row), reply_markup=task_keyboard(row["id"], row["done"]))

async def progress(message: Message):
    uid=message.from_user.id
    conn=db()
    total=conn.execute("SELECT COUNT(*) FROM tasks").fetchone()[0]
    done=conn.execute("SELECT COUNT(*) FROM completions WHERE telegram_id=?",(uid,)).fetchone()[0]
    mins=conn.execute("""
      SELECT COALESCE(SUM(t.minutes),0) FROM tasks t
      JOIN completions c ON c.task_id=t.id WHERE c.telegram_id=?
    """,(uid,)).fetchone()[0]
    conn.close()

    today=date.today().isoformat()
    rows=get_tasks(today,uid)
    td=sum(r["done"] for r in rows)
    tt=len(rows)

    await message.answer(
        f"📊 <b>IELTS Progress</b>\n\n"
        f"Today: <b>{td}/{tt}</b>\n"
        f"Overall: <b>{done}/{total}</b> ({round(done/total*100,1) if total else 0}%)\n"
        f"Study time: <b>{mins} min</b>"
    )

async def start(message: Message):
    save_user(message)
    await message.answer(
        "🎯 <b>IELTS Tracker</b>\n\n"
        "Bu bot sizning tayyor IELTS planningizni tracking qiladi.\n\n"
        "/today — bugungi tasklar\n"
        "/progress — umumiy progress\n"
        "/help — yordam"
    )

async def main():
    init_db()
    import_plan()
    bot=Bot(TOKEN)
    dp=Dispatcher()

    dp.message.register(start, CommandStart())
    dp.message.register(send_today, Command("today"))
    dp.message.register(progress, Command("progress"))

    @dp.callback_query(F.data.startswith("toggle:"))
    async def toggle(callback: CallbackQuery):
        uid=callback.from_user.id
        task_id=int(callback.data.split(":")[1])
        conn=db()
        exists=conn.execute(
            "SELECT 1 FROM completions WHERE telegram_id=? AND task_id=?",
            (uid,task_id)
        ).fetchone()
        if exists:
            conn.execute("DELETE FROM completions WHERE telegram_id=? AND task_id=?",(uid,task_id))
            done=False
        else:
            conn.execute(
                "INSERT INTO completions(telegram_id,task_id,completed_at) VALUES(?,?,?)",
                (uid,task_id,datetime.now().isoformat(timespec="seconds"))
            )
            done=True
        conn.commit(); conn.close()
        await callback.message.edit_reply_markup(reply_markup=task_keyboard(task_id,done))
        await callback.answer("✅ Bajarildi" if done else "↩️ Qaytarildi")

    @dp.message(Command("help"))
    async def help_cmd(message: Message):
        await message.answer(
            "/today — bugungi reja\n"
            "/progress — progress\n\n"
            "Task yonidagi tugmani bosib bajarilganini belgilang."
        )

    await dp.start_polling(bot)

if __name__=="__main__":
    import asyncio
    asyncio.run(main())
